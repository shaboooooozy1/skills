#!/usr/bin/env python3
"""Regression tests for bundled skill scripts, using synthetic data in evals/.

Usage: python scripts/test_skill_scripts.py   (exit 1 on any failure)
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LR = ROOT / "plugins/pc-risk-counsel/skills/analyzing-loss-runs/scripts/summarize_loss_runs.py"
RET = ROOT / "plugins/pc-risk-counsel/skills/analyzing-total-cost-of-risk/scripts/retention_analysis.py"
TCOR = ROOT / "plugins/pc-risk-counsel/skills/analyzing-total-cost-of-risk/scripts/tcor.py"
CHART = ROOT / "plugins/personal-productivity/skills/building-quick-charts/scripts/make_chart.py"
EV = ROOT / "evals"
failures = []


def run(args, cwd):
    r = subprocess.run([sys.executable, "-I", *map(str, args)], cwd=cwd, capture_output=True, text=True)
    if r.returncode:
        failures.append(f"{Path(args[0]).name} exited {r.returncode}: {r.stderr[-400:]}")
    return r


def check(cond, msg):
    if not cond:
        failures.append(msg)


with tempfile.TemporaryDirectory() as tmp:
    run([LR, EV / "analyzing-loss-runs/claims.csv", "--premium", EV / "analyzing-loss-runs/premium.csv",
         "--large-loss", "100000", "--as-of", "2026-10-24", "--out", "ls"], tmp)
    r = json.loads((Path(tmp) / "ls.json").read_text())
    t = r["totals"]
    check(t["claims"] == 21, f"claims {t['claims']} != 21")
    check(t["open_claims"] == 9, f"open {t['open_claims']} != 9")
    check(round(t["incurred"]) == 802080, f"incurred {t['incurred']} != 802080")
    check(t["claims_over_threshold"] == 3, "large-loss count != 3")
    expected_flags = {"superseded_valuation": 1, "duplicate_claim": 1, "incurred_mismatch": 1,
                      "open_zero_reserve": 2, "closed_with_reserve": 1, "loss_date_outside_term": 1,
                      "stale_valuation": 1}
    check(r["flag_counts"] == expected_flags, f"flags {r['flag_counts']} != {expected_flags}")
    y2025 = [y for y in r["by_year"] if y["policy_year"] == 2025][0]
    check(y2025["immature"] is True, "2025 should be immature")

    run([RET, EV / "analyzing-loss-runs/claims.csv", "--retentions", "25000,50000,100000",
         "--aggregate", "250000", "--lines", "AL", "--immature-years", "2024,2025", "--out", "ret"], tmp)
    s = {row["retention"]: row for row in json.loads((Path(tmp) / "ret.json").read_text())["summary"]}
    check(round(s[100000]["avg_annual_retained_all_years"]) == 61760, "100K avg all years != 61,760")
    check(round(s[100000]["avg_annual_retained_mature_years"]) == 73667, "100K avg mature != 73,667")
    check(round(s[100000]["worst_year_retained"]) == 109800, "100K worst year != 109,800")

    run([TCOR, EV / "analyzing-total-cost-of-risk/inputs.json", "--out", "tc"], tmp)
    p = json.loads((Path(tmp) / "tc.json").read_text())["periods"]
    check(round(p[0]["TCOR"]) == 531200, f"current TCOR {p[0]['TCOR']} != 531,200")
    check(round(p[1]["TCOR"]) == 533050, f"scenario TCOR {p[1]['TCOR']} != 533,050")

    (Path(tmp) / "c.csv").write_text("Year,Loss ratio\n2022,0.71\n2023,0.64\n2024,0.52\n2025,0.48\n")
    run([CHART, "c.csv", "--title", "Loss ratio fell 3 straight years", "--format", "percent",
         "--benchmark", "0.6", "--benchmark-label", "Target 60%", "--labels", "--out", "c"], tmp)
    check((Path(tmp) / "c.png").stat().st_size > 10_000, "chart PNG missing or empty")

    # Edge cases from PR #17 review: blank claim numbers, partial splits, extra PII columns,
    # duplicate premium rows, premium-only years, missing effective dates.
    tp = Path(tmp)
    (tp / "edge.csv").write_text(
        "carrier,policy_year,line,claim_number,loss_date,status,paid_loss,outstanding_loss,"
        "carrier_incurred,valuation_date,cause,claimant_name\n"
        "Test Co,2024,GL,,2024-03-01,closed,1000,0,,2026-09-30,UNREADABLE,Jane Roe\n"
        "Test Co,2024,GL,,2024-05-01,closed,2000,0,,2026-09-30,UNREADABLE,John Doe\n"
        "Test Co,2024,GL,C-3,2024-06-01,open,5000,,12000,2026-09-30,slip,Pat Poe\n"
        "Test Co,2024,GL,C-4,2024-07-01,open,,3000,,2026-09-30,slip,Sam Soe\n")
    (tp / "edge_prem.csv").write_text(
        "policy_year,line,earned_premium\n2024,GL,10000\n2024,GL,15000\n2025,GL,20000\n")
    run([LR, "edge.csv", "--premium", "edge_prem.csv", "--as-of", "2026-10-24", "--out", "edge"], tmp)
    e = json.loads((tp / "edge.json").read_text())
    check(e["totals"]["claims"] == 4, f"blank claim numbers merged: {e['totals']['claims']} claims, want 4")
    check(round(e["totals"]["incurred"]) == 1000 + 2000 + 12000 + 3000,
          f"partial split handling wrong: incurred {e['totals']['incurred']}, want 18000")
    check(e["flag_counts"].get("incomplete_split") == 2, f"incomplete_split flags {e['flag_counts']}")
    years = {y["policy_year"]: y for y in e["by_year"]}
    check(2025 in years and years[2025]["claims"] == 0, "premium-only year 2025 missing from by_year")
    yl = [r for r in e["by_year_line"] if r["policy_year"] == 2024 and r["line"] == "GL"]
    check(len(yl) == 1 and round(yl[0]["earned_premium"]) == 25000,
          f"duplicate premium rows not summed: {yl}")
    check(abs(yl[0]["loss_ratio"] - 18000 / 25000) < 1e-9 if yl else False, "2024 GL loss ratio wrong")
    check("open_claims" in e and len(e["open_claims"]) == 2, "open_claims missing from JSON")
    check(years[2024]["immature"] is True and "Dec 1" in years[2024]["maturity_basis"],
          f"maturity fallback should be conservative: {years[2024]}")
    import openpyxl
    hdr = [c.value for c in openpyxl.load_workbook(tp / "edge.xlsx")["Claims"][1]]
    check("claimant_name" not in hdr, "claimant_name leaked into Claims sheet")

    run([RET, "edge.csv", "--retentions", "10000", "--out", "edge_ret"], tmp)
    er = json.loads((tp / "edge_ret.json").read_text())
    check(round(er["summary"][0]["total_retained"]) == 1000 + 2000 + 10000 + 3000,
          f"retention replay partial/blank handling wrong: {er['summary'][0]['total_retained']}")

    # Second review round: premium coverage, written basis, undated duplicates, zero-claim years.
    (tp / "cov.csv").write_text(
        "carrier,policy_year,line,claim_number,loss_date,status,paid_loss,outstanding_loss,valuation_date\n"
        "T,2022,WC,W1,2022-05-01,closed,1000,0,2026-09-30\n"
        "T,2022,AL,A1,2022-06-01,closed,4000,0,2026-09-30\n"
        "T,2023,WC,W2,2023-05-01,closed,2000,0,2026-09-30\n"
        "T,2023,AL,A2,2023-06-01,closed,3000,0,\n"
        "T,2023,AL,A2,2023-06-01,closed,3500,0,\n"
        "T,2024,AL,A3,2024-06-01,closed,9000,0,2026-09-30\n")
    (tp / "cov_prem.csv").write_text(
        "policy_year,line,earned_premium,premium_basis\n"
        "2022,WC,10000,written\n"            # AL missing -> no 2022 year ratio
        "2023,ALL,20000,written\n"           # package covers all lines -> 2023 ratio
        "2024,AL,30000,written\n")
    run([LR, "cov.csv", "--premium", "cov_prem.csv", "--as-of", "2026-10-24", "--out", "cov"], tmp)
    c = json.loads((tp / "cov.json").read_text())
    ct, cy = c["totals"], {y["policy_year"]: y for y in c["by_year"]}
    check(cy[2022].get("loss_ratio_written") is None, f"2022 ratio should be withheld: {cy[2022]}")
    check(abs(cy[2023]["loss_ratio_written"] - 5000 / 20000) < 1e-9, f"2023 ALL ratio wrong: {cy[2023]}")
    check(ct["loss_ratio_years"] == [2023, 2024], f"rated years {ct.get('loss_ratio_years')}")
    check(abs(ct["loss_ratio_written"] - (5000 + 9000) / 50000) < 1e-9, f"all-years ratio {ct}")
    check(ct["premium_basis"] == "written" and "premium_note" in ct, "written basis not labeled")
    check(c["flag_counts"].get("premium_missing_lines") == 1, f"flags {c['flag_counts']}")
    check(ct["claims"] == 5 and c["flag_counts"].get("duplicate_claim") == 1,
          "undated duplicate not handled without crashing")

    run([RET, "cov.csv", "--retentions", "5000", "--lines", "WC", "--out", "cov_ret"], tmp)
    cr = json.loads((tp / "cov_ret.json").read_text())["summary"][0]
    check(cr["years"] == 3 and round(cr["avg_annual_retained_all_years"]) == 1000,
          f"zero-claim year not counted: {cr}")

    (tp / "blank.csv").write_text("Year,Loss ratio\n2022,0.5\n2023,\n2024,0.4\n")
    run([CHART, "blank.csv", "--title", "Gap for unknown", "--format", "percent", "--labels",
         "--out", "blank"], tmp)
    check((tp / "blank.png").exists(), "blank-cell chart failed")

    # Third review round: an old claim-free year is mature; by-line ratio uses rated years only.
    (tp / "old.csv").write_text(
        "carrier,policy_year,policy_effective,line,claim_number,loss_date,status,paid_loss,outstanding_loss,valuation_date\n"
        "T,2021,2021-07-01,GL,G1,2021-08-01,closed,5000,0,2026-09-30\n"
        "T,2023,2023-07-01,GL,G2,2023-08-01,closed,7000,0,2026-09-30\n")
    (tp / "old_prem.csv").write_text("policy_year,line,earned_premium\n2022,GL,10000\n2023,GL,10000\n")
    run([LR, "old.csv", "--premium", "old_prem.csv", "--as-of", "2026-10-24", "--out", "old"], tmp)
    o = json.loads((tp / "old.json").read_text())
    oy = {y["policy_year"]: y for y in o["by_year"]}
    check(oy[2022]["immature"] is False, f"claim-free 2022 should be mature: {oy[2022]}")
    gl = [r for r in o["by_line"] if r["line"] == "GL"][0]
    check(abs(gl["loss_ratio"] - 7000 / 20000) < 1e-9, f"by-line ratio should use rated years only: {gl}")

    # Security review: carrier text starting with '=' must stay text, never a live formula.
    (tp / "inj.csv").write_text(
        "carrier,policy_year,line,claim_number,loss_date,status,paid_loss,outstanding_loss,valuation_date,cause\n"
        'T,2024,GL,"=HYPERLINK(""http://x.test"",""ok"")",2024-06-01,closed,100,0,2026-09-30,=1+1\n')
    run([LR, "inj.csv", "--as-of", "2026-10-24", "--out", "inj"], tmp)
    run([RET, "inj.csv", "--retentions", "50", "--out", "inj_ret"], tmp)
    (tp / "inj_in.json").write_text(json.dumps({"periods": [{"label": "=2+2", "premiums": 1, "source": "=A1"}]}))
    run([TCOR, "inj_in.json", "--out", "inj_tc"], tmp)
    for book in ("inj.xlsx", "inj_ret.xlsx", "inj_tc.xlsx"):
        live = [f"{ws.title}!{c.coordinate}" for ws in openpyxl.load_workbook(tp / book).worksheets
                for row in ws.iter_rows() for c in row if c.data_type == "f"]
        check(not live, f"{book} has live formula cells: {live}")
    claims_ws = openpyxl.load_workbook(tp / "inj.xlsx")["Claims"]
    vals = [c.value for row in claims_ws.iter_rows(min_row=2) for c in row]
    check('=HYPERLINK("http://x.test","ok")' in vals, "claim number text not preserved exactly")

    (tp / "neg.csv").write_text("Item,Change\nA,500\nB,-300\nC,200\n")
    run([CHART, "neg.csv", "--title", "Credits shown", "--out", "neg"], tmp)
    check((tp / "neg.png").exists(), "negative-value chart failed")

if failures:
    print("FAILED")
    for f in failures:
        print(" -", f)
    sys.exit(1)
print("All script tests passed.")
